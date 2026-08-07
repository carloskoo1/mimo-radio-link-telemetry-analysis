from __future__ import annotations

from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches

from src.common import PROJECT_ROOT, setup_logging

def add_table(doc: Document, df: pd.DataFrame, max_rows: int = 30) -> None:
    show = df.head(max_rows)
    table = doc.add_table(rows=1, cols=len(show.columns))
    table.style = "Table Grid"
    for i, col in enumerate(show.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in show.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if pd.isna(value) else str(value)

def main() -> None:
    setup_logging("05_reportes")

    tables_dir = PROJECT_ROOT / "outputs" / "tables"
    figures_dir = PROJECT_ROOT / "outputs" / "figures"
    reports_dir = PROJECT_ROOT / "outputs" / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    tables = {}
    for path in sorted(tables_dir.glob("*.csv")):
        tables[path.stem] = pd.read_csv(path)

    with pd.ExcelWriter(reports_dir / "resultados_tesis_v3.xlsx", engine="openpyxl") as writer:
        for name, df in tables.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)

    doc = Document()
    doc.add_heading("Informe automático de resultados — Pipeline V3", 0)
    doc.add_paragraph(
        "Documento preliminar generado automáticamente. "
        "Los resultados deben revisarse antes de incorporarse al Capítulo IV."
    )

    for name, df in tables.items():
        doc.add_heading(name.replace("_", " ").title(), level=1)
        add_table(doc, df)

    for fig in sorted(figures_dir.glob("*.png")):
        doc.add_heading(fig.stem.replace("_", " ").title(), level=1)
        doc.add_picture(str(fig), width=Inches(6.2))

    doc.save(reports_dir / "informe_resultados_v3.docx")
    print("Reportes terminados.")

if __name__ == "__main__":
    main()
